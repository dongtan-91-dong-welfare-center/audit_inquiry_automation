# src/core/excel_builder.py

import os
import shutil
import openpyxl
from typing import Dict, List

class ExcelBuilder:
    """
    [엑셀 취합/생성 담당자]
    추출 및 정제가 완료된 데이터를 넘겨받아, 지정된 템플릿을 열고
    알맞은 시트의 빈 행에 데이터를 이어 붙인(Append) 뒤 저장합니다.
    """

    def __init__(self, template_dir: str = None, output_dir: str = None):
        """
        초기화 시 템플릿 폴더와 출력 폴더 경로를 설정합니다.
        입력받은 폴더 경로를 ExcelBuilder 클래스 내에서 선언 함으로써 동일 클래스 내의 다른 메서드에서도 쉽게 접근할 수 있도록 합니다.
        """
        # 현재 파일(excel_builder.py)의 절대 경로를 기준으로 프로젝트 루트 폴더를 찾음
        # src/core/excel_builder.py -> src/core -> src -> 프로젝트 루트
        current_file_path = os.path.abspath(__file__)
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(current_file_path)))
        self.template_dir = template_dir or os.path.join(project_root, "tests", "data", "templates")
        self.output_dir = output_dir or os.path.join(project_root, "tests", "data", "output")


        # 출력 폴더가 없다면 자동으로 생성합니다.
        # os.mkdir()과 다르게 os.makedirs()는 중간 폴더가 없는 경우에도 전체 경로를 생성할 수 있습니다.
        # 예를 들어 "tests/data/output" 경로가 존재하지 않을 때, os.makedirs()는 "tests", "data", "output" 폴더를 차례대로 생성합니다.
        # exist_ok=True 옵션은 이미 폴더가 존재할 때 에러를 발생시키지 않고 그냥 넘어가도록 합니다.
        os.makedirs(self.output_dir, exist_ok=True)

    def build_sheet_data_map(self, wb: openpyxl.Workbook, processed_tables: Dict[str, List[List[str]]]) -> Dict[str, List[List[str]]]:
        """
        엑셀 워크북을 열어 실제 존재하는 시트명을 확인하고,
        PostProcessor에서 넘어온 딕셔너리 데이터를 알맞은 시트명에 매핑합니다.
        """
        sheet_data_map = {}
        existing_sheets = wb.sheetnames  # 엑셀 파일에 존재하는 실제 시트명 리스트

        # 1. 금융상품 테이블 매핑 (첫 번째 시트)
        financial_data = processed_tables.get("financial_table", [])
        if financial_data:
            sheet_name = existing_sheets[0] if len(existing_sheets) > 0 else "1.금융상품"
            sheet_data_map[sheet_name] = financial_data

        # 2. 대출거래 테이블 매핑 (두 번째 시트)
        loan_data = processed_tables.get("loan_table", [])
        if loan_data:
            sheet_name = existing_sheets[1] if len(existing_sheets) > 1 else "2.대출거래"
            sheet_data_map[sheet_name] = loan_data

        return sheet_data_map

    def export_to_excel(
        self,
        company_name: str,
        bank_name: str,
        processed_tables: Dict[str, List[List[str]]],
        category: str = "은행"
    ) -> str:
        """
        데이터를 엑셀 시트에 삽입하고 파일을 저장합니다.

        Args:
            company_name: 감사대상회사명 (예: "삼성전자") - 파일명 생성에 사용
            bank_name: 조회처명 (예: "농협은행") - 각 행의 마지막 열 꼬리표로 사용
            processed_tables: 정제가 완료된 딕셔너리 표 데이터
            category: 금융권역 (기본값: "은행") - 템플릿 결정에 사용

        Returns:
            저장된 최종 엑셀 파일의 경로 (str)
        """
        company_output_dir = os.path.join(self.output_dir, company_name)
        os.makedirs(company_output_dir, exist_ok=True)
        # 1. 파일 경로 설정
        # os.path.join()은 운영체제(Windows, macOS, Linux 등)에 맞는 파일 경로 구분자(예: \ 또는 /)를 자동으로 사용하여 경로를 생성합니다.
        # 예를 들어, template_dir이 "tests/data/templates"이고 category가 "은행"일 때, template_path는 "tests/data/templates/template_은행.xlsx"가 됩니다.
        template_path = os.path.join(self.template_dir, f"template_{category}.xlsx")
        output_path = os.path.join(company_output_dir, f"{company_name}_{category}.xlsx")

        # 2. 분기 처리: 기존 파일이 없으면 템플릿 복사 (첫 회신서 처리)
        if not os.path.exists(output_path):
            if not os.path.exists(template_path):
                raise FileNotFoundError(f"❌ 템플릿 파일이 존재하지 않습니다: {template_path}\n"
                                        f"먼저 컬럼명이 적힌 템플릿을 만들어주세요.")
            # shutil은 shell utilities의 약자로, 파일 복사, 이동, 삭제 등 윈도우 탐색기에서 마우스 우클릭으로 하던
            # 파일 관리 작업들을 코드로 할 수있게 해주는 파이썬 내장 라이브러리입니다.
            # shutil 라이브러리의 copy() 함수는 template_path에 있는 파일의 내용을 output_path에 새 파일로 복사하고, 파일명은 output_path의 basename으로 지정합니다.
            shutil.copy(template_path, output_path)

        # 3. 엑셀 워크북(workbook) 열기
        # xlsxwirter는 새로운 파일을 생성할 수는 있지만, 기존 파일을 열어서 수정하는 기능은 없기 때문에, openpyxl 라이브러리를 사용
        wb = openpyxl.load_workbook(output_path)

        # 엑셀을 열고 난 뒤 실제 시트명을 읽어와서 데이터와 짝짓기(Mapping)
        sheet_data_map = self.build_sheet_data_map(wb, processed_tables)

        # 4. 시트별 데이터 삽입
        for sheet_name, data_rows in sheet_data_map.items():
            if not data_rows:
                continue  # 추출된 데이터가 없으면 해당 시트는 패스

            # 템플릿에 해당 시트가 있는지 확인
            if sheet_name not in wb.sheetnames:
                print(f"⚠️ 경고: '{sheet_name}' 시트가 템플릿에 존재하지 않아 새로 생성합니다.")
                wb.create_sheet(sheet_name)
            # worksheet
            ws = wb[sheet_name]

            # 추출된 각 행(Row)을 엑셀에 차례대로 입력
            for row in data_rows:
                # [핵심 로직] 기존 데이터 리스트의 맨 오른쪽(끝)에 '조회처' 이름을 추가
                row_to_insert = row + [bank_name]

                # ws.append()는 시트 내에서 데이터가 있는 마지막 줄을 자동으로 찾아서
                # 그 다음 빈 줄(A열부터)에 리스트의 내용을 순서대로 꽂아 넣습니다.
                ws.append(row_to_insert)

        # 5. 파일 덮어쓰기 저장 및 메모리 해제
        wb.save(output_path)
        wb.close()

        return output_path